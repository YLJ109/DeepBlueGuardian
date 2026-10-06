"""把渲染好的海报 PNG 打包成 16:9 PPTX，便于直接当封面/章节页使用。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(r"D:\front-back\OpcOcean\docs\posters")
OUT_PNG = ROOT / "out"
TARGET = ROOT / "深蓝智护-演示海报.pptx"

WIDTH = Inches(13.333)
HEIGHT = Inches(7.5)


def main() -> None:
    files = sorted(OUT_PNG.glob("*.png"))
    if not files:
        raise SystemExit("没有找到海报 PNG，请先运行 render-posters.js")

    deck = Presentation()
    deck.slide_width = WIDTH
    deck.slide_height = HEIGHT
    blank = deck.slide_layouts[6]

    for png in files:
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, width=WIDTH, height=HEIGHT)

    deck.save(str(TARGET))
    print(f"slides={len(files)}")
    print(f"file={TARGET}")


if __name__ == "__main__":
    main()
